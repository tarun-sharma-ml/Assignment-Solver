import docx
from assignment_solver.render import render_assignment
from assignment_solver.training.docx_to_blocks import parse_solved_docx

BLOCKS = [
    {"type": "paragraph", "text": "Intro paragraph."},
    {"type": "table", "rows": [["Basis", "A", "B"], ["x", "1", "2"]]},
    {"type": "heading", "text": "Limitations:"},
    {"type": "bullets", "items": ["one", "two"]},
    {"type": "heading", "text": "Python Code:"},
    {"type": "code", "text": "print('hi')\nfor i in range(2):\n    print(i)"},
    {"type": "heading", "text": "Output:"},
    {"type": "output", "text": "hi\n0\n1"},
    {"type": "example", "text": "A tiny example."},
]


def test_render_and_parse_back(tmp_path):
    out = tmp_path / "a.docx"
    render_assignment(out, student_name="Tarun Sharma", email="t@example.com", assignment_name="Demo",
                      items=[{"question": "What is X?", "blocks": BLOCKS}])
    d = docx.Document(str(out))
    assert d.paragraphs[0].text == "Name: Tarun Sharma"
    assert d.paragraphs[1].text == "Email: t@example.com"
    assert d.paragraphs[2].text == "Assignment Name: Demo"
    parsed = parse_solved_docx(out)
    assert parsed["assignment"] == "Demo"
    got = parsed["items"][0]
    assert got["question"] == "What is X?"
    assert [b["type"] for b in got["blocks"]] == [b["type"] for b in BLOCKS]
    assert got["blocks"][5]["text"] == BLOCKS[5]["text"]
