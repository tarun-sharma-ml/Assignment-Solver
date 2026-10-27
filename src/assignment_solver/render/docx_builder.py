"""Render solved blocks into a .docx that mirrors the student's template.

Template facts (measured from the sample assignments):
  * US Letter, 0.75" margins, Times New Roman
  * header: bold "Name:", "Email:", "Assignment Name:" labels
  * "Question N: ..." bold 13pt, then bold 12pt "Answer:"
  * body text 12pt; code/output in Courier New; tables with thin borders and a
    light-blue (D9E2F3) bold header row; figures centred with a caption below
"""
from __future__ import annotations

from pathlib import Path
from typing import Iterable

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Emu, Inches, Pt

BODY_PT = 12
QUESTION_PT = 13
CODE_PT = 10
HEADER_FILL = "D9E2F3"
TABLE_TWIPS = 9000
MONO = "Courier New"


def _set_font(run, size: float | None = None, bold: bool | None = None, mono: bool = False):
    if bold is not None:
        run.bold = bold
    if size:
        run.font.size = Pt(size)
    if mono:
        run.font.name = MONO
        rpr = run._element.get_or_add_rPr()
        fonts = rpr.find(qn("w:rFonts"))
        if fonts is None:
            fonts = OxmlElement("w:rFonts")
            rpr.append(fonts)
        for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            fonts.set(qn(attr), MONO)


def _para(doc, text: str = "", *, bold=False, size=BODY_PT, before=0, after=7, align=None, mono=False, style=None):
    p = doc.add_paragraph(style=style) if style else doc.add_paragraph()
    if text:
        _set_font(p.add_run(text), size=size, bold=bold or None, mono=mono)
    pf = p.paragraph_format
    pf.space_before, pf.space_after = Pt(before), Pt(after)
    if align is not None:
        p.alignment = align
    return p


def _labelled(doc, label: str, value: str, after=3):
    p = _para(doc, after=after)
    _set_font(p.add_run(label), size=BODY_PT, bold=True)
    _set_font(p.add_run(value), size=BODY_PT)
    return p


def _shade(cell, fill: str):
    tcpr = cell._element.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcpr.append(shd)


def _cell_margins(cell):
    tcpr = cell._element.get_or_add_tcPr()
    mar = OxmlElement("w:tcMar")
    for side, w in (("top", 80), ("left", 100), ("bottom", 80), ("right", 100)):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:w"), str(w))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tcpr.append(mar)


def _table_borders(table):
    tblpr = table._element.tblPr
    borders = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "auto")
        borders.append(el)
    tblpr.append(borders)


def _add_table(doc, rows: list[list[str]]):
    ncols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=ncols)
    table.autofit = False
    _table_borders(table)
    if ncols >= 3:
        first = TABLE_TWIPS // 5
        widths = [first] + [(TABLE_TWIPS - first) // (ncols - 1)] * (ncols - 1)
    else:
        widths = [TABLE_TWIPS // ncols] * ncols
    for r_idx, row in enumerate(rows):
        for c_idx in range(ncols):
            cell = table.cell(r_idx, c_idx)
            cell.width = Emu(widths[c_idx] * 635)
            _cell_margins(cell)
            para = cell.paragraphs[0]
            _set_font(para.add_run(row[c_idx] if c_idx < len(row) else ""), size=BODY_PT, bold=True if r_idx == 0 else None)
            if r_idx == 0:
                _shade(cell, HEADER_FILL)
    _para(doc, after=10)  # spacer paragraph, as in the template


def _add_bullets(doc, items: Iterable[str]):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        _set_font(p.add_run(item), size=BODY_PT)
        p.paragraph_format.space_after = Pt(4)


def _add_code(doc, text: str):
    lines = text.split("\n")
    for i, line in enumerate(lines):
        p = _para(doc, line if line else " ", size=CODE_PT, after=0 if i < len(lines) - 1 else 7, mono=True)
        p.paragraph_format.line_spacing = 1.0


def _add_figure(doc, caption: str, path: str | None):
    if path and Path(path).exists():
        p = _para(doc, before=10, after=5, align=WD_ALIGN_PARAGRAPH.CENTER)
        p.add_run().add_picture(str(path), width=Inches(5))
    _para(doc, caption, align=WD_ALIGN_PARAGRAPH.CENTER, after=10)


def _add_example(doc, text: str):
    body = text.strip()
    if body.lower().startswith("example:"):
        body = body[len("example:"):].strip()
    p = _para(doc, after=7)
    _set_font(p.add_run("Example: "), size=BODY_PT, bold=True)
    _set_font(p.add_run(body), size=BODY_PT)


def add_blocks(doc, blocks: list[dict]):
    for b in blocks:
        kind = b["type"]
        if kind == "paragraph":
            _para(doc, b["text"], after=7)
        elif kind == "heading":
            _para(doc, b["text"], bold=True, before=5, after=4)
        elif kind == "bullets":
            _add_bullets(doc, b["items"])
        elif kind == "table":
            _add_table(doc, b["rows"])
        elif kind in ("code", "output"):
            _add_code(doc, b["text"])
        elif kind == "figure":
            _add_figure(doc, b.get("caption", "Figure"), b.get("path"))
        elif kind == "example":
            _add_example(doc, b["text"])


def render_assignment(
    path: str | Path,
    *,
    student_name: str,
    email: str,
    assignment_name: str,
    items: list[dict],
) -> Path:
    """items: [{"question": str, "blocks": [block, ...]}, ...] (numbered automatically)."""
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, side, Inches(0.75))
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(BODY_PT)
    rpr = normal.element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.append(fonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        fonts.set(qn(attr), "Times New Roman")

    _labelled(doc, "Name: ", student_name)
    _labelled(doc, "Email: ", email)
    _labelled(doc, "Assignment Name: ", assignment_name, after=10)

    for n, item in enumerate(items, start=1):
        _para(doc, f"Question {n}: {item['question']}", bold=True, size=QUESTION_PT, before=10 if n > 1 else 4, after=5)
        _para(doc, "Answer:", bold=True, after=4)
        add_blocks(doc, item["blocks"])

    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))
    return path
