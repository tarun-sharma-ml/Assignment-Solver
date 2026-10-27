"""Turn one of the student's *solved* .docx files back into (question, blocks) pairs.

This is how the training data is made: the rendered template is read back into
the same block JSON the model is asked to produce.
"""
from __future__ import annotations

import re
from pathlib import Path

import docx
from docx.table import Table
from docx.text.paragraph import Paragraph

from ..ingest.docx_reader import iter_block_items

_QUESTION = re.compile(r"^\s*Question\s*\d+\s*[:.)\-–]?\s*", re.I)
_MONO = {"courier new", "consolas", "courier", "lucida console"}


def _is_mono(p: Paragraph) -> bool:
    runs = [r for r in p.runs if r.text.strip()]
    return bool(runs) and all((r.font.name or "").lower() in _MONO for r in runs)


def _is_bold(p: Paragraph) -> bool:
    runs = [r for r in p.runs if r.text.strip()]
    return bool(runs) and all(r.bold for r in runs)


def _has_drawing(p: Paragraph) -> bool:
    return bool(p._p.xpath(".//w:drawing"))


def _is_list(p: Paragraph) -> bool:
    name = (p.style.name if p.style is not None else "").lower()
    return name.startswith("list")


class _Builder:
    def __init__(self):
        self.blocks: list[dict] = []
        self._mono: list[str] = []
        self._bullets: list[str] = []
        self._last_heading = ""
        self._want_caption = False

    def _flush_mono(self):
        if self._mono:
            kind = "output" if self._last_heading.lower().startswith("output") else "code"
            self.blocks.append({"type": kind, "text": "\n".join(self._mono).strip("\n")})
            self._mono = []

    def flush(self):
        self._flush_mono()
        if self._bullets:
            self.blocks.append({"type": "bullets", "items": self._bullets})
            self._bullets = []

    def add_paragraph(self, p: Paragraph):
        text = p.text.strip()
        if _has_drawing(p):
            self.flush()
            self.blocks.append({"type": "figure", "caption": "Figure"})
            self._want_caption = True
            return
        if not text:
            return
        if _is_mono(p):
            if self._bullets:
                self.flush()
            self._mono.append(p.text.rstrip("\n"))
            return
        if _is_list(p):
            self._want_caption = False
            self._flush_mono()
            self._bullets.append(text)
            return
        self.flush()
        if self._want_caption and text.lower().startswith("figure"):
            self.blocks[-1]["caption"] = text
            self._want_caption = False
            return
        self._want_caption = False
        if text.lower().startswith("example:"):
            self.blocks.append({"type": "example", "text": text[len("example:"):].strip()})
        elif _is_bold(p) and len(text) <= 80:
            self._last_heading = text
            self.blocks.append({"type": "heading", "text": text})
        else:
            self.blocks.append({"type": "paragraph", "text": text})

    def add_table(self, t: Table):
        self.flush()
        rows = [[c.text.strip() for c in row.cells] for row in t.rows]
        if rows:
            self.blocks.append({"type": "table", "rows": rows})

    def finish(self) -> list[dict]:
        self.flush()
        return self.blocks


def parse_solved_docx(path: str | Path) -> dict:
    """Return {"assignment": str, "items": [{"question": str, "blocks": [...]}, ...]}."""
    document = docx.Document(str(path))
    assignment = Path(path).stem
    items: list[dict] = []
    current: dict | None = None
    builder: _Builder | None = None

    def close():
        nonlocal current, builder
        if current is not None and builder is not None:
            current["blocks"] = builder.finish()
            if current["blocks"]:
                items.append(current)
        current = builder = None

    for item in iter_block_items(document):
        if isinstance(item, Paragraph):
            text = item.text
            if current is None and re.match(r"^\s*Assignment Name", text, re.I):
                assignment = text.split(":", 1)[-1].strip() or assignment
                continue
            if _QUESTION.match(text):
                close()
                lines = [l for l in text.split("\n") if l.strip()]
                question = _QUESTION.sub("", lines[0], count=1).strip()
                current, builder = {"question": question, "blocks": []}, _Builder()
                continue
            if current is None:
                continue  # header lines (Name / Email)
            if text.strip().lower() == "answer:":
                continue
            builder.add_paragraph(item)
        elif current is not None:
            builder.add_table(item)
    close()
    return {"assignment": assignment, "items": items}
